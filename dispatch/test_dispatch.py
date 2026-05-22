"""Smoke tests for the Truer Foods dispatch tool."""
from pathlib import Path

from openpyxl import load_workbook

import dispatch
import zones

SAMPLE = Path(__file__).parent / "sample_orders.csv"


def test_extract_fsa():
    assert zones.extract_fsa("V6X 2C9") == "V6X"
    assert zones.extract_fsa("V6X2C9") == "V6X"
    assert zones.extract_fsa("1133 W Broadway Vancouver BC V6H 1G1") == "V6H"
    assert zones.extract_fsa("V6X") == "V6X"
    assert zones.extract_fsa("") is None
    assert zones.extract_fsa("no postal here") is None


def test_zone_for_known_areas():
    assert zones.zone_for("V6X 2C9")[0] == "Richmond"
    assert zones.zone_for("V5C 2H8")[0] == "Burnaby"
    assert zones.zone_for("V7P 2R5")[0] == "North/West Van"
    assert zones.zone_for("V4A 2H9")[0] == "White Rock"


def test_zone_for_unknown():
    assert zones.zone_for("Z9Z 9Z9") == ("Out of Area", 99)
    assert zones.zone_for("") == ("Out of Area", 99)


def test_load_sample():
    stops = dispatch.load_stops(SAMPLE)
    assert len(stops) == 14
    assert any(s.customer == "Tojo's Restaurant" for s in stops)


def test_sort_is_by_drive_order():
    stops = dispatch.sort_stops(dispatch.load_stops(SAMPLE))
    orders = [s.drive_order for s in stops]
    assert orders == sorted(orders)
    # First stops should be Richmond, last should be Surrey/WhiteRock area.
    assert stops[0].zone == "Richmond"


def test_split_drivers_balances_and_preserves_stops():
    stops = dispatch.sort_stops(dispatch.load_stops(SAMPLE))
    for n in (1, 2, 3, 4):
        routes = dispatch.split_drivers(stops, n)
        assert sum(len(r) for r in routes) == len(stops)
        assert len(routes) <= n
        for r in routes:
            assert r, "should not produce empty routes"


def test_split_keeps_zones_contiguous():
    stops = dispatch.sort_stops(dispatch.load_stops(SAMPLE))
    routes = dispatch.split_drivers(stops, 3)
    for route in routes:
        # Zones inside a route should appear in non-decreasing drive_order
        orders = [s.drive_order for s in route]
        assert orders == sorted(orders)


def test_end_to_end(tmp_path):
    out = tmp_path / "routes.xlsx"
    rc = dispatch.main([str(SAMPLE), "--drivers", "2", "--out", str(out)])
    assert rc == 0
    assert out.exists()
    wb = load_workbook(out)
    assert "Summary" in wb.sheetnames
    assert "Driver 1" in wb.sheetnames
    assert "Driver 2" in wb.sheetnames
    # Driver 1 should have header row + data rows
    d1 = wb["Driver 1"]
    assert d1.max_row > 4
