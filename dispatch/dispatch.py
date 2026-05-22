"""Daily route dispatch for Truer Foods (Richmond, BC).

Reads a CSV of stops, groups them by Vancouver-area zones in a sensible
drive order, splits them across N drivers, and writes an Excel workbook
with one route sheet per driver plus a summary tab.

Usage:
    python dispatch.py orders.csv --drivers 3 --out today.xlsx
"""
from __future__ import annotations

import argparse
import math
import sys
import urllib.parse
from dataclasses import dataclass
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from zones import HOME_BASE, extract_fsa, zone_for

REQUIRED_COLS = {"customer", "address"}


@dataclass
class Stop:
    customer: str
    address: str
    postal: str
    zone: str
    drive_order: int
    boxes: int
    product_type: str
    window: str
    notes: str

    @property
    def maps_url(self) -> str:
        return "https://www.google.com/maps/search/?api=1&query=" + urllib.parse.quote_plus(self.address)


def load_stops(csv_path: Path) -> list[Stop]:
    df = pd.read_csv(csv_path, dtype=str).fillna("")
    df.columns = [c.strip().lower() for c in df.columns]
    missing = REQUIRED_COLS - set(df.columns)
    if missing:
        raise SystemExit(f"orders CSV is missing required columns: {sorted(missing)}")

    stops: list[Stop] = []
    for _, row in df.iterrows():
        customer = row["customer"].strip()
        if not customer:
            continue
        address = row["address"].strip()
        postal_field = (row.get("postal_code", "") or "").strip().upper()
        fsa = extract_fsa(postal_field) or extract_fsa(address) or ""
        zone, order = zone_for(postal_field or address)
        try:
            boxes = int((row.get("boxes", "") or "1").strip() or "1")
        except ValueError:
            boxes = 1
        stops.append(Stop(
            customer=customer,
            address=address,
            postal=fsa,
            zone=zone,
            drive_order=order,
            boxes=boxes,
            product_type=(row.get("product_type", "") or "").strip(),
            window=(row.get("window", "") or "").strip(),
            notes=(row.get("notes", "") or "").strip(),
        ))
    return stops


def sort_stops(stops: list[Stop]) -> list[Stop]:
    """Order by drive sequence, then by postal so neighbouring streets stay together."""
    return sorted(stops, key=lambda s: (s.drive_order, s.postal, s.customer))


def split_drivers(stops: list[Stop], n_drivers: int) -> list[list[Stop]]:
    """Split sorted stops into N driver routes, keeping zones contiguous.

    Walks the sorted list and advances to the next driver once the current
    driver hits the target stop count AND the zone is about to change. This
    avoids splitting a single zone across two trucks.
    """
    if n_drivers <= 1 or len(stops) <= 1:
        return [stops] if stops else []
    target = math.ceil(len(stops) / n_drivers)
    routes: list[list[Stop]] = [[] for _ in range(n_drivers)]
    i = 0
    for stop in stops:
        if i < n_drivers - 1 and len(routes[i]) >= target and routes[i] and routes[i][-1].zone != stop.zone:
            i += 1
        routes[i].append(stop)
    return [r for r in routes if r]


def write_excel(routes: list[list[Stop]], out_path: Path) -> None:
    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="305496")

    summary = wb.active
    summary.title = "Summary"
    summary.append(["Truer Foods Daily Dispatch"])
    summary["A1"].font = Font(bold=True, size=14)
    summary.append([f"Home base: {HOME_BASE}"])
    summary.append([])
    summary.append(["Driver", "Stops", "Boxes", "Zones"])
    for cell in summary[4]:
        cell.font = header_font
        cell.fill = header_fill
    for i, route in enumerate(routes, start=1):
        zones_in_route: list[str] = []
        for s in route:
            if s.zone not in zones_in_route:
                zones_in_route.append(s.zone)
        summary.append([
            f"Driver {i}",
            len(route),
            sum(s.boxes for s in route),
            ", ".join(zones_in_route),
        ])
    for col, width in enumerate([12, 8, 8, 60], start=1):
        summary.column_dimensions[get_column_letter(col)].width = width

    for i, route in enumerate(routes, start=1):
        sheet = wb.create_sheet(f"Driver {i}")
        sheet.append([f"Driver {i} — Start: {HOME_BASE}"])
        sheet["A1"].font = Font(bold=True, size=12)
        sheet.append([f"{len(route)} stops · {sum(s.boxes for s in route)} boxes total"])
        sheet.append([])
        headers = ["#", "Zone", "Customer", "Address", "Postal",
                   "Boxes", "Product", "Window", "Notes", "Map"]
        sheet.append(headers)
        for cell in sheet[4]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        for n, stop in enumerate(route, start=1):
            row_idx = sheet.max_row + 1
            sheet.append([
                n,
                stop.zone,
                stop.customer,
                stop.address,
                stop.postal,
                stop.boxes,
                stop.product_type,
                stop.window,
                stop.notes,
                "Open in Maps",
            ])
            link_cell = sheet.cell(row=row_idx, column=10)
            link_cell.hyperlink = stop.maps_url
            link_cell.font = Font(color="0563C1", underline="single")

        for col, width in enumerate([4, 18, 26, 40, 8, 6, 10, 14, 28, 14], start=1):
            sheet.column_dimensions[get_column_letter(col)].width = width
        sheet.freeze_panes = "A5"

    wb.save(out_path)


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(description="Truer Foods route dispatcher")
    p.add_argument("orders_csv", type=Path, help="Path to today's orders CSV")
    p.add_argument("--drivers", type=int, default=1, help="Number of drivers (default 1)")
    p.add_argument("--out", type=Path, default=Path("routes.xlsx"),
                   help="Output Excel path (default routes.xlsx)")
    args = p.parse_args(argv)

    if args.drivers < 1:
        raise SystemExit("--drivers must be >= 1")

    stops = load_stops(args.orders_csv)
    if not stops:
        raise SystemExit("No stops found in orders CSV")

    routes = split_drivers(sort_stops(stops), args.drivers)
    write_excel(routes, args.out)

    zones_used = sorted({s.zone for s in stops})
    unknown = sum(1 for s in stops if s.zone == "Out of Area")
    print(f"Wrote {args.out} — {len(stops)} stops across {len(routes)} driver(s).")
    print(f"Zones: {', '.join(zones_used)}")
    if unknown:
        print(f"WARNING: {unknown} stop(s) had no recognized postal code "
              "(bucketed in 'Out of Area' at the end of the last route). "
              "Add a postal_code column entry or fix the address.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
